"""Bounded, value-only XLSX/CSV lists; exports never execute imported formulas."""

import csv
import io
import json
import re
import zipfile
from datetime import date, datetime, time
from pathlib import PurePosixPath

from defusedxml.ElementTree import iterparse
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.http import HttpResponse
from django.utils.translation import gettext_lazy as _
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

MAX_LINKS = 200
MAX_COLUMNS = 20
MAX_UPLOAD_BYTES = 1024 * 1024
MAX_DATA_BYTES = 256 * 1024
MAX_CELL_CHARS = 1000


def checked_xlsx(data):
    """Bound work before openpyxl reads shared strings, styles or worksheet XML.

    Nothing is extracted to disk. Every XML member passes a DTD/entity-free,
    depth/node-limited parser, including members not selected for import.
    """
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 64 or len({e.filename for e in entries}) != len(entries):
            raise ValueError("archive entries")
        total, nodes = 0, 0
        for entry in entries:
            name = entry.filename
            if (
                entry.flag_bits & 1
                or entry.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED)
                or name.startswith("/")
                or "\\" in name
                or ".." in PurePosixPath(name).parts
                or not name.endswith((".xml", ".rels", "/"))
                or name.startswith(("xl/externalLinks/", "xl/embeddings/"))
            ):
                raise ValueError("unsupported archive member")
            total += entry.file_size
            if entry.file_size > 4 * 1024**2 or total > 8 * 1024**2:
                raise ValueError("expanded archive size")
            if entry.is_dir():
                continue
            with archive.open(entry) as source:
                xml = source.read(4 * 1024**2 + 1)
            if len(xml) != entry.file_size or len(xml) > 4 * 1024**2:
                raise ValueError("expanded member size")
            depth = 0
            for event, element in iterparse(
                io.BytesIO(xml), events=("start", "end"), forbid_dtd=True
            ):
                tag = element.tag.rsplit("}", 1)[-1]
                if event == "start":
                    depth += 1
                    nodes += 1
                    if depth > 32 or nodes > 50000:
                        raise ValueError("XML complexity")
                    if name.startswith("xl/worksheets/") and tag == "c":
                        reference = element.get("r", "")
                        match = re.fullmatch(
                            r"([A-Z]{1,3})([1-9][0-9]{0,6})", reference
                        )
                        if (
                            not match
                            or column_index_from_string(match[1]) > MAX_COLUMNS
                            or int(match[2]) > MAX_LINKS + 1
                        ):
                            raise ValidationError(
                                _(
                                    "La liste est limitée à 200 lignes et 20 colonnes, en plus de l’en-tête."
                                )
                            )
                    if tag == "f" and name.startswith("xl/worksheets/"):
                        raise ValidationError(
                            _(
                                "Le classeur contient des formules. Collez leurs résultats en valeurs avant l’import."
                            )
                        )
                else:
                    depth -= 1
                    element.clear()


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        value = value.isoformat()
    elif isinstance(value, bool):
        value = "TRUE" if value else "FALSE"
    elif isinstance(value, float) and value.is_integer():
        value = str(int(value))
    else:
        value = str(value)
    # XML 1.0 controls are rejected for consistent HTML, CSV and XLSX output.
    if len(value) > MAX_CELL_CHARS or re.search(
        r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", value
    ):
        raise ValidationError(
            _(
                "Chaque cellule doit contenir au plus 1 000 caractères, sans caractères de contrôle."
            )
        )
    return value


def normalized(rows):
    rows = [[cell_text(value) for value in row] for row in rows]
    rows = [row for row in rows if any(value.strip() for value in row)]
    if len(rows) < 2:
        raise ValidationError(
            _("Ajoutez une ligne d’en-tête et au moins une ligne de données.")
        )
    if len(rows) > MAX_LINKS + 1:
        raise ValidationError(
            _("La liste est limitée à 200 lignes et 20 colonnes, en plus de l’en-tête.")
        )
    width = max(
        (i + 1 for row in rows for i, value in enumerate(row) if value != ""), default=0
    )
    if width > MAX_COLUMNS:
        raise ValidationError(
            _("La liste est limitée à 200 lignes et 20 colonnes, en plus de l’en-tête.")
        )
    rows = [(row + [""] * width)[:width] for row in rows]
    headers = [value.strip() for value in rows[0]]
    if any(not value or len(value) > 80 for value in headers) or len(
        {h.casefold() for h in headers}
    ) != len(headers):
        raise ValidationError(
            _(
                "Les en-têtes doivent être uniques, non vides et limités à 80 caractères."
            )
        )
    if len(json.dumps(rows, ensure_ascii=False).encode()) > MAX_DATA_BYTES:
        raise ValidationError(
            _(
                "Le contenu de la liste dépasse 256 Kio. Réduisez les colonnes ou le nombre de lignes."
            )
        )
    return headers, rows[1:]


def read_list(upload):
    from .storage import scan

    suffix = PurePosixPath(upload.name).suffix.lower()
    if suffix not in (".xlsx", ".csv"):
        raise ValidationError(_("Choisissez un fichier Excel .xlsx ou CSV UTF-8."))
    data = upload.read(MAX_UPLOAD_BYTES + 1)
    if upload.size > MAX_UPLOAD_BYTES or len(data) > MAX_UPLOAD_BYTES:
        raise ValidationError(_("La liste ne doit pas dépasser 1 Mio."))
    book = None
    try:
        if suffix == ".xlsx":
            checked_xlsx(data)
        scan(SimpleUploadedFile("list" + suffix, data))
        if suffix == ".xlsx":
            book = load_workbook(
                io.BytesIO(data), read_only=True, data_only=False, keep_links=False
            )
            if not book.worksheets:
                raise ValueError("no worksheet")
            sheet = book.worksheets[0]
            sheet.reset_dimensions()  # Do not trust the supplied dimension range.
            rows = list(
                sheet.iter_rows(
                    max_row=MAX_LINKS + 1, max_col=MAX_COLUMNS, values_only=True
                )
            )
        else:
            text = data.decode("utf-8-sig")
            try:
                delimiter = (
                    csv.Sniffer().sniff(text[:8192], delimiters=",;\t").delimiter
                )
            except csv.Error:
                delimiter = ";" if ";" in text.splitlines()[0] else ","
            rows = []
            for row in csv.reader(
                io.StringIO(text, newline=""), delimiter=delimiter, strict=True
            ):
                if len(row) > MAX_COLUMNS:
                    raise ValidationError(
                        _(
                            "La liste est limitée à 200 lignes et 20 colonnes, en plus de l’en-tête."
                        )
                    )
                if any(value.strip() for value in row):
                    rows.append(row)
                if len(rows) > MAX_LINKS + 1:
                    raise ValidationError(
                        _(
                            "La liste est limitée à 200 lignes et 20 colonnes, en plus de l’en-tête."
                        )
                    )
        return normalized(rows)
    except ValidationError:
        raise
    except Exception:
        # Third-party parser errors never expose workbook contents or internals.
        raise ValidationError(
            _(
                "Liste illisible ou format non pris en charge. Utilisez le modèle Excel ou un CSV UTF-8 avec uniquement des valeurs."
            )
        ) from None
    finally:
        if book is not None:
            book.close()


def csv_text(value):
    value = str(value)
    if re.sub(r"^[\s\ufeff]+", "", value).startswith(
        ("=", "+", "-", "@")
    ) or value.startswith(("\t", "\r", "\n")):
        return "'" + value
    return value


def spreadsheet_response(
    headers, rows, format, *, filename="linkvault-links", url_column=None
):
    """Create a new plain table, with no imported formulas or external content."""
    if format == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output, delimiter=";", quoting=csv.QUOTE_ALL)
        writer.writerow([csv_text(value) for value in headers])
        writer.writerows([csv_text(value) for value in row] for row in rows)
        response = HttpResponse(
            "\ufeff" + output.getvalue(), content_type="text/csv; charset=utf-8"
        )
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = "LinkVault"
        for row_number, row in enumerate([headers, *rows], 1):
            for column, value in enumerate(row, 1):
                cell = sheet.cell(row_number, column)
                cell.value = str(value)
                cell.data_type = "s"  # Even =, +, - and @ prefixes remain text.
                cell.number_format = "@"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_number == 1:
                    cell.fill = PatternFill("solid", fgColor="245E48")
                    cell.font = Font(color="FFFFFF", bold=True)
                elif column - 1 == url_column:
                    cell.hyperlink = str(
                        value
                    )  # Only server-generated capability URLs.
                    cell.font = Font(color="245E48", underline="single")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 30
        for column in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(column)].width = (
                60 if column - 1 == url_column else 26
            )
        output = io.BytesIO()
        book.save(output)
        book.close()
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    response["Content-Disposition"] = f'attachment; filename="{filename}.{format}"'
    return response
