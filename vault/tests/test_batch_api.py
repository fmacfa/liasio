import csv
import io
import json
import secrets
import threading
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from unittest.mock import patch

from django.contrib.auth.hashers import check_password, make_password
from django.db import close_old_connections, connections
from django.test import Client, TestCase, TransactionTestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from vault.api_auth import credential_version, digest
from vault.mfa import cipher
from vault.models import (
    APIReceipt,
    APIToken,
    Audit,
    Collection,
    Event,
    LinkBatch,
    LinkBatchItem,
    RateBucket,
    ShareLink,
    Webhook,
    WebhookDelivery,
)
from .test_security import IsolatedFiles
from .test_capabilities import legacy_capability


def api_key(owner, scopes=None):
    token = APIToken(
        owner=owner,
        name="Batch workflow",
        scopes=scopes or ["read", "links:write"],
        credential_version=credential_version(owner),
        expires_at=timezone.now() + timedelta(days=1),
    )
    raw = f"lv_{token.pk.hex}.{secrets.token_urlsafe(32)}"
    token.digest = digest(raw)
    token.save()
    return token, raw


class BatchAPITests(IsolatedFiles, TestCase):
    def setUp(self):
        super().setUp()
        self.user = self.account()
        self.folder = self.collection(self.user)
        self.token, self.raw = api_key(self.user)
        self.url = reverse("api_link_batch_create", args=[self.folder.pk])

    def post(self, data=None, key="batch-request-001", raw=None, url=None):
        return self.client.post(
            url or self.url,
            data=json.dumps(
                data if data is not None else {"name": "Clients", "names": ["A", "B"]}
            ),
            content_type="application/json",
            HTTP_AUTHORIZATION="Bearer " + (raw or self.raw),
            HTTP_IDEMPOTENCY_KEY=key,
        )

    def get(self, url, raw=None):
        return self.client.get(url, HTTP_AUTHORIZATION="Bearer " + (raw or self.raw))

    def test_named_creation_returns_ordered_urls_and_is_visible_in_browser(self):
        values = {
            "name": "Contacts",
            "names": ["A", "Zoé", "A"],
            "conditions": {
                "expires_in_days": 3,
                "max_downloads": 1,
                "password": " a ",
                "allow_previews": True,
            },
        }
        with patch("vault.api_batches.make_password", wraps=make_password) as hasher:
            response = self.post(values)
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(hasher.call_count, 1)
        data = response.json()
        self.assertEqual([row["label"] for row in data["items"]], values["names"])
        self.assertEqual([row["position"] for row in data["items"]], [1, 2, 3])
        self.assertEqual(len({row["url"] for row in data["items"]}), 3)
        self.assertEqual((data["state"], data["count"]), ("complete", 3))
        self.assertEqual(data["source_rows"], [[name] for name in values["names"]])
        self.assertEqual(Event.objects.filter(kind="created").count(), 3)
        self.assertEqual(Audit.objects.filter(action="link_created").count(), 3)
        self.assertEqual(LinkBatchItem.objects.count(), 3)
        for link in ShareLink.objects.all():
            self.assertTrue(check_password(" a ", link.password_hash))
            self.assertEqual(link.max_downloads, 1)
            self.assertTrue(link.allow_previews)
            self.assertLess(
                abs((link.expires_at - timezone.now()).total_seconds() - 3 * 86400), 5
            )
        self.client.force_login(self.user)
        self.assertContains(self.client.get(reverse("link_batches")), "Contacts")
        self.assertContains(
            self.client.get(reverse("link_batch", args=[data["id"]])), 'value="Zoé"'
        )
        read = self.get(reverse("api_link_batch", args=[data["id"]]))
        self.assertEqual(read.json(), data)
        self.assertEqual(read["Cache-Control"], "no-store, private")
        self.assertEqual(read["X-Content-Type-Options"], "nosniff")
        self.assertNotIn(" a ", APIReceipt.objects.get().response)
        self.assertNotIn(data["items"][0]["url"], APIReceipt.objects.get().response)

    def test_upgrade_preserves_api_replays_and_exports_current_compact_urls(self):
        # A pre-upgrade receipt must replay verbatim, even though fresh reads use
        # the new URL representation for those same database links.
        with patch("vault.security.capability", side_effect=legacy_capability):
            original = self.post()
        self.assertEqual(original.status_code, 201)
        old = original.json()
        self.assertEqual(self.post().json(), old)
        current = self.get(reverse("api_link_batch", args=[old["id"]])).json()
        urls = [row["url"] for row in current["items"]]
        self.assertEqual(
            [row["id"] for row in current["items"]],
            [row["id"] for row in old["items"]],
        )
        for row, old_row in zip(current["items"], old["items"], strict=True):
            self.assertRegex(row["url"], r"/s/[A-Za-z0-9_-]{43}/\Z")
            self.assertRegex(old_row["url"], r"/s/[0-9a-f]{32}\.[0-9a-f]{64}/\Z")
            self.assertEqual(self.client.get(row["url"]).status_code, 200)
            self.assertEqual(self.client.get(old_row["url"]).status_code, 200)
        fresh = self.post(key="compact-creation-request").json()
        for row in fresh["items"]:
            self.assertRegex(row["url"], r"/s/[A-Za-z0-9_-]{43}/\Z")
        csv_response = self.get(
            reverse("api_link_batch_export", args=[old["id"], "csv"])
        )
        self.assertEqual(csv_response.status_code, 200)
        rows = list(csv.DictReader(
            io.StringIO(csv_response.content.decode("utf-8-sig")), delimiter=";"
        ))
        self.assertEqual([row["Liasio URL"] for row in rows], urls)
        xlsx_response = self.get(
            reverse("api_link_batch_export", args=[old["id"], "xlsx"])
        )
        self.assertEqual(xlsx_response.status_code, 200)
        book = load_workbook(io.BytesIO(xlsx_response.content), read_only=True)
        rows = list(book.active.values)
        column = rows[0].index("Liasio URL")
        self.assertEqual([row[column] for row in rows[1:]], urls)
        book.close()
        self.client.force_login(self.user)
        self.assertContains(self.client.get(reverse("link_batches")), old["name"])
        detail = self.client.get(reverse("link_batch", args=[old["id"]]))
        for url in urls:
            self.assertContains(detail, url)
            self.assertContains(self.client.get(reverse("links")), url)

    def test_quantity_limits_and_optional_conditions(self):
        response = self.post({"name": "Generated", "count": 200})
        self.assertEqual(response.status_code, 201, response.content)
        items = response.json()["items"]
        self.assertEqual(len(items), 200)
        self.assertEqual(items[0]["label"], "Link 001")
        self.assertEqual(items[-1]["label"], "Link 200")
        self.assertIsNone(items[0]["expires_at"])
        self.assertIsNone(items[0]["max_downloads"])
        self.assertFalse(items[0]["password_protected"])
        self.assertFalse(items[0]["allow_previews"])

    def test_replay_is_identical_after_expiry_and_quota_without_hashing_again(self):
        now = timezone.now()
        values = {
            "name": "Retry",
            "names": ["A"],
            "conditions": {
                "expires_at": (now + timedelta(seconds=20)).isoformat(),
                "password": "a",
            },
        }
        created = self.post(values)
        self.assertEqual(created.status_code, 201)
        with (
            patch(
                "vault.api_batches.timezone.now",
                return_value=now + timedelta(minutes=1),
            ),
            patch("vault.batches.MAX_SAVED_BATCHES", 1),
            patch("vault.batches.rate_limited", return_value=True),
            patch("vault.api_batches.make_password") as hasher,
        ):
            replayed = self.post(values)
        self.assertEqual(replayed.status_code, 201, replayed.content)
        self.assertEqual(replayed.json(), created.json())
        self.assertEqual(replayed["Idempotency-Replayed"], "true")
        hasher.assert_not_called()
        conflict = self.post({**values, "names": ["B"]})
        self.assertEqual(conflict.status_code, 409)
        self.assertEqual(conflict.json()["error"]["code"], "idempotency_conflict")
        self.assertEqual(LinkBatch.objects.count(), 1)
        self.assertEqual(ShareLink.objects.count(), 1)
        receipt = APIReceipt.objects.get()
        payload = json.loads(cipher().decrypt(receipt.response.encode()))
        self.assertEqual(payload, created.json())
        self.assertNotIn('"password"', json.dumps(payload))

    def test_invalid_names_types_and_conditions_leave_no_partial_batches(self):
        invalid = [
            {},
            {"name": "No names"},
            {"name": 7, "count": 1},
            {"name": "", "count": 1},
            {"name": "x" * 81, "count": 1},
            {"name": "Bad", "names": [], "count": 1},
            *[{"name": "Bad", "count": value} for value in (0, 201, True, "2", 1.5)],
            *[
                {"name": "Bad", "names": value}
                for value in (
                    "A",
                    [],
                    ["A"] * 201,
                    [1],
                    [None],
                    [{}],
                    [""],
                    [" "],
                    ["x" * 81],
                    ["bad\nname"],
                    ["\ud800"],
                    ["\ufffe"],
                )
            ],
            {"name": "Bad", "names": ["A"], "owner_id": str(self.user.pk)},
            *[
                {"name": "Bad", "names": ["A"], "conditions": value}
                for value in (
                    None,
                    [],
                    {"unknown": 1},
                    {"expiry_mode": "none"},
                    {"expires_in_days": True},
                    {"expires_in_days": "2"},
                    {"expires_in_days": 0},
                    {"max_downloads": True},
                    {"max_downloads": -1},
                    {"max_downloads": 100001},
                    {"allow_previews": "true"},
                    {"password": 1},
                    {"password": "x" * 129},
                    {"expires_at": "2030-01-01T12:00:00"},
                    {"expires_at": "garbage"},
                    {"expires_at": "2000-01-01T12:00:00Z"},
                    {"expires_at": "2030-01-01T12:00:00Z", "expires_in_days": 2},
                )
            ],
        ]
        for value in invalid:
            with self.subTest(value=repr(value)):
                RateBucket.objects.all().delete()
                response = self.post(value)
                self.assertEqual(response.status_code, 400, response.content)
                self.assertEqual(response.json()["error"]["code"], "invalid_request")
                self.assertFalse(LinkBatch.objects.exists())
                self.assertFalse(ShareLink.objects.exists())
                self.assertFalse(APIReceipt.objects.exists())
        self.assertEqual(self.post(key="").status_code, 400)

    def test_body_boundary_and_bad_json_fail_before_creation(self):
        with patch("vault.api_batches.inputs") as parser:
            response = self.client.generic(
                "POST", self.url, b"x", CONTENT_LENGTH=str(256 * 1024 + 1)
            )
        self.assertEqual(response.status_code, 413)
        parser.assert_not_called()
        for body in (
            b"not json",
            b"[]",
            b"{",
            b'{"names":' + b"[" * 1500 + b"]" * 1500 + b"}",
        ):
            response = self.client.post(
                self.url,
                body,
                content_type="application/json",
                HTTP_AUTHORIZATION="Bearer " + self.raw,
            )
            self.assertEqual(response.status_code, 400)
        self.assertFalse(LinkBatch.objects.exists())

    def test_bearer_scope_and_token_lifecycle(self):
        self.client.force_login(self.user)
        self.assertEqual(
            self.client.post(self.url, {"name": "Cookie", "count": 1}).status_code, 401
        )
        self.assertEqual(
            self.client.get(
                reverse("api_link_batches"), {"api_key": self.raw}
            ).status_code,
            401,
        )
        _, read_only = api_key(self.user, ["read"])
        self.assertEqual(self.post(raw=read_only).status_code, 403)
        _, write_only = api_key(self.user, ["links:write"])
        created = self.post(raw=write_only)
        self.assertEqual(created.status_code, 201)
        batch_id = created.json()["id"]
        for path in (
            reverse("api_link_batches"),
            reverse("api_link_batch", args=[batch_id]),
            reverse("api_link_batch_export", args=[batch_id, "csv"]),
        ):
            self.assertEqual(self.get(path, write_only).status_code, 403)
            self.assertEqual(self.get(path, read_only).status_code, 200)
        self.token.revoked_at = timezone.now()
        self.token.save()
        self.assertEqual(self.post().status_code, 401)
        self.token.revoked_at = None
        self.token.expires_at = timezone.now() - timedelta(seconds=1)
        self.token.save()
        self.assertEqual(self.post().status_code, 401)
        self.token.expires_at = timezone.now() + timedelta(days=1)
        self.token.save()
        self.user.mfa_version = uuid.uuid4()
        self.user.save()
        self.assertEqual(self.post().status_code, 401)
        self.assertEqual(LinkBatch.objects.count(), 1)

    def test_owner_trash_and_deleted_batch_boundaries(self):
        data = self.post().json()
        other = self.account("other@example.com")
        _, raw = api_key(other)
        for path in (
            reverse("api_link_batch", args=[data["id"]]),
            reverse("api_link_batch_export", args=[data["id"], "xlsx"]),
        ):
            self.assertEqual(self.get(path, raw).status_code, 404)
        self.assertEqual(self.post(raw=raw).status_code, 404)
        self.assertEqual(
            self.get(reverse("api_link_batches"), raw).json()["results"], []
        )
        self.folder.deleted_at = timezone.now()
        self.folder.save()
        self.assertEqual(self.post(key="after-trash-request").status_code, 404)
        self.assertEqual(
            self.get(reverse("api_link_batch", args=[data["id"]])).status_code, 404
        )
        self.assertEqual(self.get(reverse("api_link_batches")).json()["results"], [])
        self.folder.deleted_at = None
        self.folder.save()
        LinkBatch.objects.get(pk=data["id"]).delete()
        self.assertEqual(
            self.get(reverse("api_link_batch", args=[data["id"]])).status_code, 404
        )
        self.assertEqual(ShareLink.objects.count(), 2)

    def test_browser_batch_and_draft_consultation_with_pagination(self):
        draft = LinkBatch.objects.create(
            collection=self.folder,
            name="Browser draft",
            headers=["Email"],
            source_rows=[["example@example.com"]],
        )
        detail = self.get(reverse("api_link_batch", args=[draft.pk])).json()
        self.assertEqual(detail["state"], "draft")
        self.assertEqual(detail["items"], [])
        self.assertIsNone(detail["exports"])
        self.assertEqual(
            self.get(
                reverse("api_link_batch_export", args=[draft.pk, "csv"])
            ).status_code,
            404,
        )
        LinkBatch.objects.filter(pk=draft.pk).update(
            created_at=timezone.now() - timedelta(days=2)
        )
        self.assertEqual(
            self.get(reverse("api_link_batch", args=[draft.pk])).json()["state"],
            "expired",
        )
        LinkBatch.objects.bulk_create(
            [
                LinkBatch(collection=self.folder, name=f"Draft {n}", source_rows=[[]])
                for n in range(50)
            ]
        )
        first = self.get(reverse("api_link_batches")).json()
        self.assertEqual(len(first["results"]), 50)
        self.assertEqual(first["next_offset"], 50)
        last = self.get(reverse("api_link_batches") + "?offset=50").json()
        self.assertEqual(len(last["results"]), 1)
        self.assertIsNone(last["next_offset"])
        self.assertNotIn("source_rows", first["results"][0])
        empty = Collection.objects.create(owner=self.user, title="Other")
        filtered = self.get(reverse("api_link_batches") + f"?collection_id={empty.pk}")
        self.assertEqual(filtered.json()["results"], [])
        self.assertEqual(
            self.get(reverse("api_link_batches") + "?offset=-1").status_code, 400
        )
        self.assertEqual(
            self.get(reverse("api_link_batches") + "?collection_id=bad").status_code,
            400,
        )

    def test_exports_preserve_pairing_and_current_labels_as_inert_values(self):
        data = self.post({"name": "Export", "names": ["=1+1", "00127", "Zoé"]}).json()
        link = ShareLink.objects.get(pk=data["items"][0]["id"])
        link.label = "@SUM(1)"
        link.save()
        response = self.get(reverse("api_link_batch_export", args=[data["id"], "xlsx"]))
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment;", response["Content-Disposition"])
        self.assertEqual(response["Cache-Control"], "no-store, private")
        book = load_workbook(io.BytesIO(response.content))
        sheet = book.active
        self.assertEqual(sheet["A2"].value, "=1+1")
        self.assertEqual(sheet["A2"].data_type, "s")
        self.assertEqual(sheet["B2"].value, "@SUM(1)")
        self.assertEqual(sheet["B2"].data_type, "s")
        self.assertEqual(sheet["C2"].hyperlink.target, data["items"][0]["url"])
        self.assertEqual(sheet["A3"].value, "00127")
        book.close()
        csv_response = self.get(
            reverse("api_link_batch_export", args=[data["id"], "csv"])
        )
        rows = list(
            csv.reader(
                io.StringIO(csv_response.content.decode("utf-8-sig")), delimiter=";"
            )
        )
        self.assertEqual(rows[1][:2], ["'=1+1", "'@SUM(1)"])
        self.assertEqual(
            [row[2] for row in rows[1:]], [item["url"] for item in data["items"]]
        )
        current = self.get(reverse("api_link_batch", args=[data["id"]])).json()
        self.assertEqual(current["items"][0]["label"], "@SUM(1)")
        self.assertEqual(current["items"][0]["url"], data["items"][0]["url"])
        self.assertEqual(
            self.get(
                reverse("api_link_batch_export", args=[data["id"], "exe"])
            ).status_code,
            404,
        )

    def test_scan_quota_and_rate_limits_are_shared_with_browser(self):
        with self.settings(ANTIVIRUS_REQUIRED=True):
            self.assertEqual(self.post().status_code, 400)
        self.assertFalse(LinkBatch.objects.exists())
        self.assertFalse(ShareLink.objects.exists())
        empty = Collection.objects.create(owner=self.user, title="Empty")
        self.assertEqual(
            self.post(
                url=reverse("api_link_batch_create", args=[empty.pk])
            ).status_code,
            400,
        )
        created = self.post()
        self.assertEqual(created.status_code, 201)
        with (
            patch("vault.batches.MAX_SAVED_BATCHES", 1),
            patch("vault.api_batches.make_password") as hasher,
        ):
            self.assertEqual(self.post(key="quota-request-002").status_code, 400)
        hasher.assert_not_called()
        with patch("vault.batches.rate_limited", return_value=True):
            limited = self.post(key="rate-request-002")
            self.assertEqual(limited.status_code, 429)
            self.assertEqual(limited["Retry-After"], "60")
            exported = self.get(
                reverse("api_link_batch_export", args=[created.json()["id"], "csv"])
            )
            self.assertEqual(exported.status_code, 429)
        # The browser budget is the API budget too, including different API keys.
        self.client.force_login(self.user)
        RateBucket.objects.all().delete()
        browser_url = reverse("link_batch_prepare", args=[self.folder.pk])
        for _ in range(10):
            self.assertEqual(
                self.client.post(
                    browser_url, {"name": "Draft", "mode": "count", "count": 1}
                ).status_code,
                302,
            )
        self.assertEqual(self.post(key="shared-rate-request").status_code, 429)

    def test_database_failure_rolls_back_list_links_receipt_and_webhooks(self):
        hook = Webhook.objects.create(
            owner=self.user,
            name="Hook",
            destination=cipher().encrypt(b"https://example.com/hook").decode(),
            secret=cipher().encrypt(b"test").decode(),
            events=["link.created"],
        )
        original = Event.objects.create
        calls = []

        def fail(*args, **kwargs):
            calls.append(1)
            if len(calls) == 2:
                raise RuntimeError("simulated failure")
            return original(*args, **kwargs)

        with patch("vault.batches.Event.objects.create", side_effect=fail):
            with self.assertRaises(RuntimeError):
                self.post()
        for model in (LinkBatch, ShareLink, LinkBatchItem, APIReceipt, WebhookDelivery):
            self.assertFalse(model.objects.exists(), model.__name__)
        self.assertEqual(Event.objects.filter(kind="created").count(), 0)
        created = self.post()
        self.assertEqual(created.status_code, 201)
        self.post()
        self.assertEqual(hook.deliveries.count(), 2)
        for delivery in hook.deliveries.all():
            self.assertEqual(
                set(delivery.payload["data"]), {"collection_id", "link_id"}
            )
            self.assertNotIn('"url"', json.dumps(delivery.payload))

    def test_authorization_and_time_rechecked_after_password_hashing(self):
        def revoke(password):
            APIToken.objects.filter(pk=self.token.pk).update(revoked_at=timezone.now())
            return make_password(password)

        with patch("vault.api_batches.make_password", side_effect=revoke):
            response = self.post(
                {"name": "Race", "names": ["A"], "conditions": {"password": "a"}}
            )
        self.assertEqual(response.status_code, 404)
        self.assertFalse(LinkBatch.objects.exists())
        APIToken.objects.filter(pk=self.token.pk).update(revoked_at=None)
        future = timezone.now() + timedelta(minutes=1)
        values = {
            "name": "Expire",
            "names": ["A"],
            "conditions": {"expires_at": future.isoformat(), "password": "a"},
        }

        def expire(password):
            clock.start()
            return make_password(password)

        clock = patch(
            "vault.api_batches.timezone.now", return_value=future + timedelta(seconds=1)
        )
        try:
            with patch("vault.api_batches.make_password", side_effect=expire):
                self.assertEqual(self.post(values).status_code, 400)
        finally:
            clock.stop()
        self.assertFalse(LinkBatch.objects.exists())


class ConcurrentBatchAPITests(IsolatedFiles, TransactionTestCase):
    def test_simultaneous_api_retry_hashes_without_lock_then_commits_one_batch(self):
        owner = self.account()
        folder = self.collection(owner)
        token, raw = api_key(owner)
        barrier = threading.Barrier(2)
        payload = json.dumps(
            {"name": "Race", "names": ["A", "B"], "conditions": {"password": "a"}}
        )

        def hash_without_owner_lock(value):
            self.assertFalse(connections["default"].in_atomic_block)
            barrier.wait(timeout=10)
            return make_password(value)

        def submit(_):
            close_old_connections()
            try:
                response = Client().post(
                    reverse("api_link_batch_create", args=[folder.pk]),
                    payload,
                    content_type="application/json",
                    HTTP_AUTHORIZATION="Bearer " + raw,
                    HTTP_IDEMPOTENCY_KEY="concurrent-api-batch",
                )
                return response.status_code, response.json()
            finally:
                connections.close_all()

        with patch(
            "vault.api_batches.make_password", side_effect=hash_without_owner_lock
        ):
            with ThreadPoolExecutor(max_workers=2) as pool:
                responses = list(pool.map(submit, range(2)))
        self.assertEqual([status for status, _ in responses], [201, 201])
        self.assertEqual(responses[0][1], responses[1][1])
        self.assertEqual(LinkBatch.objects.count(), 1)
        self.assertEqual(ShareLink.objects.count(), 2)
        self.assertEqual(Event.objects.filter(kind="created").count(), 2)
        self.assertEqual(APIReceipt.objects.filter(token=token).count(), 1)
