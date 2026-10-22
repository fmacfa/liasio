import csv
import io
import json
import threading
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from unittest.mock import patch

from django.contrib.auth.hashers import check_password, make_password
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import close_old_connections, connection, connections
from django.test import Client, TestCase, TransactionTestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from vault import batches, link_lists
from vault.models import (
    Audit,
    Event,
    LinkBatch,
    LinkBatchItem,
    ShareLink,
    Webhook,
    WebhookDelivery,
)
from vault.mfa import cipher
from .test_security import IsolatedFiles


def workbook(rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def rewrite_xlsx(data, replace=None, additions=None):
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as destination,
    ):
        for name in source.namelist():
            content = source.read(name)
            if replace:
                content = replace(name, content)
            destination.writestr(name, content)
        for name, content in (additions or {}).items():
            destination.writestr(name, content)
    return output.getvalue()


class BatchListTests(IsolatedFiles, TestCase):
    def test_excel_csv_preserve_order_extra_columns_and_text_identifiers(self):
        rows = [
            ["Référence", "Nom", "Email"],
            ["00127", "Zoé", "zoe@example.com"],
            ["00128", "Alex", "alex@example.com"],
        ]
        for name, data in (
            ("list.xlsx", workbook(rows)),
            (
                "list.csv",
                "\ufeffRéférence;Nom;Email\r\n00127;Zoé;zoe@example.com\r\n00128;Alex;alex@example.com\r\n".encode(),
            ),
        ):
            with self.subTest(name=name):
                self.assertEqual(
                    link_lists.read_list(SimpleUploadedFile(name, data)),
                    (rows[0], rows[1:]),
                )
        self.assertEqual(
            link_lists.read_list(
                SimpleUploadedFile(
                    "list.csv", b'Name,Note\r\nAlex,"first line\nsecond line"\r\n'
                )
            ),
            (["Name", "Note"], [["Alex", "first line\nsecond line"]]),
        )

    def test_rejects_bad_headers_sizes_formats_formula_workbooks_and_macros(self):
        invalid = [
            ("list.xls", b"old binary"),
            ("list.csv", b"\xff\xfe"),
            ("list.csv", b"Name\n"),
            ("list.csv", b"Name,Name\na,b\n"),
            ("list.csv", b"Name,\na,b\n"),
            ("list.csv", b"Name\n" + b"a\n" * 201),
            ("list.csv", ("Name\n" + "x" * 1001).encode()),
            ("list.csv", b"Name\na\x00b"),
            ("list.csv", b"x" * (1024**2 + 1)),
            (
                "list.xlsx",
                workbook([["Name"], ['=HYPERLINK("https://example.com","A")']]),
            ),
            (
                "list.xlsx",
                rewrite_xlsx(
                    workbook([["Name"], ["A"]]),
                    additions={"xl/vbaProject.bin": b"macro"},
                ),
            ),
            ("list.xlsx", workbook([[str(n) for n in range(21)], list(range(21))])),
        ]
        for name, data in invalid:
            with self.subTest(name=name, size=len(data)):
                with self.assertRaises(ValidationError):
                    link_lists.read_list(SimpleUploadedFile(name, data))

    def test_archive_bombs_entities_depth_duplicates_and_forged_dimensions(self):
        original = workbook([["Name"], ["Camille"]])
        attack_members = [
            {"docProps/big.xml": b" " * (4 * 1024**2 + 1)},
            {
                "docProps/entities.xml": b'<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><x>&e;</x>'
            },
            {"docProps/deep.xml": b"<x>" * 40 + b"</x>" * 40},
            {"xl/externalLinks/externalLink1.xml": b"<x/>"},
            {"../unsafe.xml": b"<x/>"},
        ]
        for members in attack_members:
            with self.subTest(members=list(members)):
                with self.assertRaises(ValidationError):
                    link_lists.read_list(
                        SimpleUploadedFile(
                            "bad.xlsx", rewrite_xlsx(original, additions=members)
                        )
                    )
        with self.assertWarns(UserWarning):
            repeated = rewrite_xlsx(original, additions={"xl/workbook.xml": b"<x/>"})
        with self.assertRaises(ValidationError):
            link_lists.read_list(SimpleUploadedFile("bad.xlsx", repeated))
        forged = rewrite_xlsx(
            original,
            replace=lambda name, data: (
                data.replace(b'ref="A1:A2"', b'ref="A1:XFD1048576"')
                if name.endswith("sheet1.xml")
                else data
            ),
        )
        self.assertEqual(
            link_lists.read_list(SimpleUploadedFile("safe.xlsx", forged)),
            (["Name"], [["Camille"]]),
        )
        outside = rewrite_xlsx(
            original,
            replace=lambda name, data: (
                data.replace(b'r="A2"', b'r="A999999"')
                if name.endswith("sheet1.xml")
                else data
            ),
        )
        with self.assertRaises(ValidationError):
            link_lists.read_list(SimpleUploadedFile("bad.xlsx", outside))

    def test_antivirus_is_required_when_enabled_and_rejection_saves_nothing(self):
        with self.settings(ANTIVIRUS_REQUIRED=True, CLAMAV_HOST=""):
            with self.assertRaisesMessage(ValidationError, "antivirus"):
                link_lists.read_list(SimpleUploadedFile("list.csv", b"Name\nCamille"))
        with patch("vault.storage.scan", side_effect=ValidationError("blocked")):
            with self.assertRaisesMessage(ValidationError, "blocked"):
                link_lists.read_list(
                    SimpleUploadedFile("list.xlsx", workbook([["Name"], ["Camille"]]))
                )

    def test_exports_treat_formula_prefixes_as_text_and_keep_zeroes(self):
        values = [
            "=1+1",
            "+SUM(1,2)",
            "-1+2",
            "@SUM(A1)",
            " \t=1",
            "\tDDE",
            "\u00a0=1+1",
            "00123",
            "Zoé",
        ]
        response = link_lists.spreadsheet_response(
            ["Name"], [[value] for value in values], "xlsx"
        )
        book = load_workbook(io.BytesIO(response.content))
        for number, value in enumerate(values, 2):
            cell = book.active.cell(number, 1)
            self.assertEqual(cell.value, value)
            self.assertEqual(cell.data_type, "s")
        book.close()
        response = link_lists.spreadsheet_response(
            ["=Header"], [[value] for value in values], "csv"
        )
        rows = list(
            csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";")
        )
        self.assertEqual(rows[0][0], "'=Header")
        self.assertTrue(all(row[0].startswith("'") for row in rows[1:8]))
        self.assertEqual(rows[-2][0], "00123")


class BulkLinkTests(IsolatedFiles, TestCase):
    def setUp(self):
        super().setUp()
        self.user = self.account()
        self.folder = self.collection(self.user)
        self.client.force_login(self.user)

    def prepare(self, count=3):
        response = self.client.post(
            reverse("link_batch_prepare", args=[self.folder.pk]),
            {"name": "Clients", "mode": "count", "count": count},
        )
        self.assertEqual(response.status_code, 302)
        return LinkBatch.objects.latest("created_at")

    def create(self, batch, **changes):
        data = {
            "action": "create",
            "expiry_mode": "days",
            "expires_in_days": 5,
            "max_downloads": 1,
            "password": " a ",
            "allow_previews": "on",
        }
        data.update(
            {f"label_{i}": f"Client {i + 1}" for i in range(len(batch.source_rows))}
        )
        data.update(changes)
        return self.client.post(reverse("link_batch", args=[batch.pk]), data)

    def test_quantity_create_rename_copy_export_and_repeat_do_not_duplicate(self):
        batch = self.prepare()
        with patch("vault.batches.make_password", wraps=make_password) as hasher:
            response = self.create(batch)
        self.assertEqual(response.status_code, 302)
        self.assertEqual(hasher.call_count, 1)
        links = [item.link for item in batch.items.select_related("link")]
        self.assertEqual(len({link.public_url for link in links}), 3)
        self.assertEqual(Event.objects.filter(kind="created").count(), 3)
        self.assertEqual(Audit.objects.filter(action="link_created").count(), 3)
        for link in links:
            self.assertTrue(check_password(" a ", link.password_hash))
            self.assertEqual(link.max_downloads, 1)
            self.assertEqual((link.expires_at - timezone.now()).days, 4)
        self.assertEqual(self.create(batch).status_code, 302)
        self.assertEqual(ShareLink.objects.count(), 3)
        original = links[0].public_url
        response = self.client.post(
            reverse("link_rename", args=[links[0].pk]), {"label_0": "Camille"}
        )
        self.assertEqual(response.status_code, 302)
        links[0].refresh_from_db()
        self.assertEqual((links[0].label, links[0].public_url), ("Camille", original))
        response = self.client.post(
            reverse("link_batch", args=[batch.pk]),
            {
                "action": "rename",
                "label_0": "Renamed",
                "label_1": "Alex",
                "label_2": "Zoé",
            },
        )
        self.assertEqual(response.status_code, 302)
        exported = self.client.get(
            reverse("link_batch_export", args=[batch.pk, "xlsx"])
        )
        self.assertEqual(exported.status_code, 200)
        book = load_workbook(io.BytesIO(exported.content))
        self.assertEqual(book.active["A2"].value, "Renamed")
        self.assertEqual(book.active["B2"].value, original)
        self.assertEqual(book.active["B2"].hyperlink.target, original)
        self.assertNotIn(
            "password", " ".join(str(cell.value).lower() for cell in book.active[1])
        )
        book.close()

    def test_import_mapping_and_export_keep_source_order_and_metadata(self):
        source = [
            ["Référence", "Email", "Nom"],
            ["00042", "zoe@example.com", "Zoé"],
            ["00043", "alex@example.com", "Alex"],
        ]
        response = self.client.post(
            reverse("link_batch_prepare", args=[self.folder.pk]),
            {
                "name": "Import",
                "mode": "file",
                "source": SimpleUploadedFile("list.xlsx", workbook(source)),
            },
        )
        self.assertEqual(response.status_code, 302)
        batch = LinkBatch.objects.get()
        response = self.client.get(response.url)
        self.assertEqual(response.context["column"], 2)
        self.assertContains(response, 'value="Zoé"')
        self.assertEqual(
            self.create(batch, label_0="Zoé", label_1="Alex").status_code, 302
        )
        response = self.client.get(reverse("link_batch_export", args=[batch.pk, "csv"]))
        rows = list(
            csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";")
        )
        self.assertEqual(rows[0][:3], source[0])
        self.assertEqual([row[:3] for row in rows[1:]], source[1:])
        self.assertEqual([row[3] for row in rows[1:]], ["Zoé", "Alex"])
        self.assertTrue(
            all(row[4].startswith("http") and "/s/" in row[4] for row in rows[1:])
        )

    def test_single_use_is_independent_for_each_link(self):
        batch = self.prepare(2)
        self.assertEqual(self.create(batch, password="").status_code, 302)
        links = [item.link for item in batch.items.select_related("link")]
        from vault.security import capability

        url = "/s/" + capability("share", links[0].pk) + "/download/"
        response = self.client.post(url)
        self.assertEqual(response.status_code, 200)
        b"".join(response.streaming_content)
        for link in links:
            link.refresh_from_db()
        self.assertEqual([link.downloads for link in links], [1, 0])
        self.assertEqual([link.status_code for link in links], ["exhausted", "active"])

    def test_labels_and_imported_values_are_escaped_and_exports_are_private(self):
        batch = LinkBatch.objects.create(
            collection=self.folder,
            name="<script>batch</script>",
            headers=["Name", "Note"],
            source_rows=[
                ["<img src=x onerror=alert(1)>", "</script><script>alert(1)</script>"]
            ],
        )
        response = self.client.get(reverse("link_batch", args=[batch.pk]))
        self.assertNotContains(response, "<script>alert(1)</script>")
        self.assertContains(response, "\\u003C/script\\u003E")
        self.create(batch)
        response = self.client.get(
            reverse("link_batch_export", args=[batch.pk, "xlsx"])
        )
        self.assertEqual(response["Cache-Control"], "no-store, private")
        self.assertEqual(response["X-Content-Type-Options"], "nosniff")
        self.client.logout()
        for name, args in (
            ("link_batch", [batch.pk]),
            ("link_batch_export", [batch.pk, "csv"]),
            ("link_batch_template", ["xlsx"]),
        ):
            self.assertEqual(self.client.get(reverse(name, args=args)).status_code, 302)

    def test_invalid_conditions_names_and_scan_state_create_nothing(self):
        batch = self.prepare()
        for change in (
            {"expires_in_days": 0},
            {"label_1": ""},
            {"label_2": "x" * 81},
            {"label_0": "bad\x0bname"},
            {"max_downloads": -1},
        ):
            with self.subTest(change=change):
                self.assertEqual(self.create(batch, **change).status_code, 200)
                self.assertEqual(ShareLink.objects.count(), 0)
        with self.settings(ANTIVIRUS_REQUIRED=True):
            response = self.create(batch)
            self.assertContains(response, "antivirus")
            self.assertEqual(ShareLink.objects.count(), 0)
        self.folder.files.update(deleted_at=timezone.now())
        self.assertContains(self.create(batch), "Ajoutez au moins un fichier")
        self.assertFalse(LinkBatchItem.objects.exists())

    def test_transaction_rolls_back_entire_batch_and_events_on_failure(self):
        batch = self.prepare()
        original = Event.objects.create
        calls = []

        def fails_after_first(*args, **kwargs):
            calls.append(1)
            if len(calls) == 2:
                raise RuntimeError("simulated database failure")
            return original(*args, **kwargs)

        with patch("vault.batches.Event.objects.create", side_effect=fails_after_first):
            with self.assertRaises(RuntimeError):
                self.create(batch)
        self.assertFalse(ShareLink.objects.exists())
        self.assertFalse(LinkBatchItem.objects.exists())
        self.assertFalse(Event.objects.exists())
        batch.refresh_from_db()
        self.assertIsNone(batch.completed_at)

    def test_no_cross_account_access_or_trash_access_and_csrf_is_required(self):
        batch = self.prepare()
        self.create(batch)
        link = batch.items.first().link
        owner_client = Client(enforce_csrf_checks=True)
        owner_client.force_login(self.user)
        self.assertEqual(
            owner_client.post(
                reverse("link_batch", args=[batch.pk]),
                {"action": "rename", "label_0": "Forged"},
            ).status_code,
            403,
        )
        self.assertEqual(
            owner_client.post(
                reverse("link_batch_prepare", args=[self.folder.pk]),
                {"name": "Forged", "mode": "count", "count": 2},
            ).status_code,
            403,
        )
        self.client.force_login(self.account("other@example.com"))
        paths = [
            reverse("link_batch", args=[batch.pk]),
            reverse("link_batch_export", args=[batch.pk, "xlsx"]),
            reverse("link_batch_delete", args=[batch.pk]),
            reverse("link_rename", args=[link.pk]),
            reverse("link_batch_prepare", args=[self.folder.pk]),
        ]
        for url in paths:
            with self.subTest(url=url):
                self.assertEqual(
                    self.client.post(url).status_code, 405 if "/export/" in url else 404
                )
        self.assertEqual(self.client.get(paths[1]).status_code, 404)
        self.client.force_login(self.user)
        self.folder.deleted_at = timezone.now()
        self.folder.save()
        self.assertEqual(self.client.get(paths[0]).status_code, 404)
        self.assertEqual(self.client.get(paths[1]).status_code, 404)
        self.assertEqual(self.client.post(paths[3], {"label_0": "No"}).status_code, 404)

    def test_200_links_supported_with_field_limits_and_bad_counts_rejected(self):
        url = reverse("link_batch_prepare", args=[self.folder.pk])
        for count in ("0", "201", "1.5", "1e2", "-1"):
            self.assertEqual(
                self.client.post(
                    url, {"name": "Invalid", "mode": "count", "count": count}
                ).status_code,
                200,
            )
        self.assertFalse(LinkBatch.objects.exists())
        batch = self.prepare(200)
        self.assertEqual(self.create(batch).status_code, 302)
        self.assertEqual(ShareLink.objects.count(), 200)
        response = self.client.get(reverse("link_batch", args=[batch.pk]))
        self.assertContains(response, 'name="label_199"')

    def test_request_limits_and_account_rate_limit_apply_before_parsing(self):
        url = reverse("link_batch_prepare", args=[self.folder.pk])
        self.assertEqual(
            self.client.generic(
                "POST", url, b"x", CONTENT_LENGTH=str(1152 * 1024 + 1)
            ).status_code,
            413,
        )
        self.assertEqual(
            self.client.generic("POST", url, b"", CONTENT_LENGTH="").status_code, 411
        )
        with (
            patch("vault.security.rate_allowed", return_value=False),
            patch("vault.batches.read_list") as parser,
        ):
            response = self.client.post(
                url,
                {
                    "name": "No",
                    "mode": "file",
                    "source": SimpleUploadedFile("x.csv", b"Name\nA"),
                },
            )
        self.assertEqual(response.status_code, 429)
        parser.assert_not_called()
        with patch("vault.batches.MAX_SAVED_BATCHES", 1):
            self.prepare(1)
            response = self.client.post(
                url, {"name": "No", "mode": "count", "count": 1}
            )
            self.assertEqual(response.status_code, 200)
            self.assertEqual(LinkBatch.objects.count(), 1)

    def test_delete_list_keeps_links_and_expired_drafts_are_cleaned(self):
        batch = self.prepare()
        self.create(batch)
        ids = list(ShareLink.objects.values_list("pk", flat=True))
        self.assertEqual(
            self.client.post(reverse("link_batch_delete", args=[batch.pk])).status_code,
            302,
        )
        self.assertFalse(LinkBatch.objects.exists())
        self.assertFalse(LinkBatchItem.objects.exists())
        self.assertCountEqual(ShareLink.objects.values_list("pk", flat=True), ids)
        old = self.prepare()
        LinkBatch.objects.filter(pk=old.pk).update(
            created_at=timezone.now() - timedelta(days=2)
        )
        self.assertContains(self.create(old), "brouillon a expiré")
        batches.cleanup_drafts()
        self.assertFalse(LinkBatch.objects.exists())
        self.assertEqual(ShareLink.objects.count(), 3)

    def test_events_enqueue_once_per_link_with_no_urls_passwords_or_imported_emails(
        self,
    ):
        key = "dGVzdC1iYXRjaC1rZXktMzItYnl0ZXMtbG9uZy0wMDA="
        with self.settings(MFA_ENCRYPTION_KEY=key):
            hook = Webhook.objects.create(
                owner=self.user,
                name="Test",
                destination=cipher().encrypt(b"https://example.com/hook").decode(),
                secret=cipher().encrypt(b"test-secret").decode(),
                events=["link.created"],
            )
            batch = self.prepare(2)
            self.create(batch)
            self.create(batch)
            deliveries = WebhookDelivery.objects.filter(webhook=hook)
            self.assertEqual(deliveries.count(), 2)
            for delivery in deliveries:
                self.assertEqual(
                    set(delivery.payload["data"]), {"collection_id", "link_id"}
                )
                self.assertNotIn("password", json.dumps(delivery.payload))


class ConcurrentBatchTests(IsolatedFiles, TransactionTestCase):
    def test_simultaneous_submit_produces_one_batch(self):
        user = self.account()
        folder = self.collection(user)
        batch = LinkBatch.objects.create(
            collection=folder, name="Race", source_rows=[[], []]
        )
        barrier = threading.Barrier(2)
        clients = []
        for _number in range(2):
            client = Client()
            client.force_login(user)
            clients.append(client)

        def submit(client):
            close_old_connections()
            try:
                barrier.wait(timeout=10)
                return client.post(
                    reverse("link_batch", args=[batch.pk]),
                    {
                        "action": "create",
                        "label_0": "A",
                        "label_1": "B",
                        "expiry_mode": "none",
                    },
                ).status_code
            finally:
                connections.close_all()

        with ThreadPoolExecutor(max_workers=2) as pool:
            codes = list(pool.map(submit, clients))
        self.assertEqual(codes, [302, 302], connection.vendor)
        self.assertEqual(ShareLink.objects.count(), 2)
        self.assertEqual(Event.objects.filter(kind="created").count(), 2)
        self.assertEqual(LinkBatchItem.objects.count(), 2)
